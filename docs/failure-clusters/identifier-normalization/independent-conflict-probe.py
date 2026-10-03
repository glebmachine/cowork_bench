import importlib.util, subprocess, tempfile, pathlib, json, hashlib, zipfile, io, datetime
import openpyxl
root=pathlib.Path('/private/tmp/cowork-fix-identifier-normalization')
s=importlib.util.spec_from_file_location('test',root/'tests/identifier_normalization/test_native_checkers.py'); t=importlib.util.module_from_spec(s);s.loader.exec_module(t)
a=t.NativeCheckers()
def mutate(wb):
 t.change('coupon',wb,True)
 ws=wb['Data_Analysis']; col=[c.value for c in ws[1]].index('Market_Avg_Price')+1; extra=ws.max_column+1
 ws.cell(1,extra).value='Benchmark_Market_Avg_Price'
 for row in range(2,ws.max_row+1):
  ws.cell(row,extra).value=ws.cell(row,col).value
  ws.cell(row,col).value=999999
for variant in ('candidate','baseline'):
 old=t.load
 if variant=='baseline':
  def load(case):
   source=subprocess.check_output(['git','show',f'd943e75:tasks/finalpool/{t.CASES[case]}/evaluation/main.py'],cwd=root)
   import types
   mod=types.ModuleType('baseline');exec(compile(source,'baseline/main.py','exec'),mod.__dict__); return mod
  t.load=load
 r=a.evaluate('coupon',mutate)
 print(variant,{k:v for k,v in r.items() if k.startswith(('Data_Analysis Category','Market_Avg_Price','Price_Gap_Pct'))})
 t.load=old
prov=json.loads((root/'docs/failure-clusters/identifier-normalization/provenance.json').read_text())
with zipfile.ZipFile('/Users/glebmikheev/prj/ouroboros-bank-backend/artifacts/evals/cowork-full496-20261003/raw-runs.zip') as z:
 for case,p in prov.items():
  raw=z.read(p['archive_member']); assert hashlib.sha256(raw).hexdigest()==p['xlsx_sha256']
  assert hashlib.sha256(z.read(p['baseline_eval_member'])).hexdigest()==p['baseline_eval_sha256']
  wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=True)
  sheets={ws.title:[[str(v) if isinstance(v,(datetime.datetime,datetime.date,datetime.time)) else v for v in row] for row in ws.iter_rows(values_only=True)] for ws in wb}
  fixture=json.loads((root/f'tests/identifier_normalization/fixtures/{case}.json').read_text())
  assert sheets==fixture['sheets'],case
  print('verified xlsx, eval SHA and complete cell projection:',case)
