"""Native-checker harness; archived agent code is read only, never executed."""
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

from docx import Document

ROOT = Path(os.environ.get("GRADER_REPO_ROOT", Path(__file__).resolve().parents[2]))
FIXTURES = Path(__file__).parent / "fixtures/remaining-language"


def grader(case):
    spec = importlib.util.spec_from_file_location(case, ROOT / "tasks/finalpool" / case / "evaluation/main.py")
    module = importlib.util.module_from_spec(spec)
    with patch.dict(sys.modules, {"psycopg2": types.ModuleType("psycopg2")}):
        spec.loader.exec_module(module)
    checks = {}
    module.check = module.record = lambda name, ok, detail="", **kwargs: checks.__setitem__(name, bool(ok))
    return module, checks


class Harness(unittest.TestCase):
    def noise(self, extra=None, replace=None, table=False, other=None):
        module, checks = grader("arxiv-latex-review-teamly-word")
        document = Document(FIXTURES / "LLM_Paper_Synthesis.docx")
        if replace is not None:
            for paragraph in document.paragraphs:
                if 'robot learning' in paragraph.text.lower():
                    paragraph.text = replace
        if extra:
            document.add_paragraph(extra)
        if table:
            document.add_table(rows=1, cols=1).cell(0, 0).text = 'Robot Learning with Affordances: method and results.'
        if other:
            module._NOISE_CORPUS.update(other)
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            document.save(Path(tmp) / "LLM_Paper_Synthesis.docx")
            module.check_word(tmp)
            module.check_noise_exclusion()
        return checks["Robotics noise paper is correctly EXCLUDED from all deliverables"]

    def formula(self, source=None):
        module, checks = grader("terminal-insales-pdf-excel-word-gform")
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            path = Path(tmp) / "quality_audit.py"
            if source is None:
                shutil.copyfile(FIXTURES / "quality_audit.py.txt", path)
            else:
                path.write_text(source)
            module.check_script(tmp)
        return checks["CRITICAL: quality_audit.py references the score formula"]

