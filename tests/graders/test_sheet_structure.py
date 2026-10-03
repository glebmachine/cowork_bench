"""Offline regression tests against the actual, standalone task evaluators."""
import contextlib
import importlib.util
import io
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SUMMARY_CASES = (
    'canvas-exam-prep-scheduler', 'canvas-grade-summary', 'canvas-quiz-report',
    'canvas-ta-workload-excel-email', 'sf-hr-manager-report', 'sf-hr-salary-growth',
)


def grader(case):
    spec = importlib.util.spec_from_file_location(case, ROOT / 'tasks/finalpool' / case / 'evaluation/main.py')
    module = importlib.util.module_from_spec(spec)
    # No database queries are used by these Excel tests. Fail closed if one appears.
    stub = types.ModuleType('psycopg2')
    def no_db(**kwargs):
        raise RuntimeError('DB unavailable in offline Excel regression')
    stub.connect = no_db
    with patch.dict(sys.modules, {'psycopg2': stub}):
        spec.loader.exec_module(module)
    return module


def book(sheet, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    return wb


class SummaryTests(unittest.TestCase):
    def test_header_is_optional_without_losing_first_metric(self):
        for case in SUMMARY_CASES:
            with self.subTest(case=case):
                g = grader(case)
                data = [['Total_Courses', 22], ['Overall_Avg_Score', 71.25]]
                for rows in (data, [['Metric', 'Value']] + data):
                    loaded = g.load_sheet_rows(book('Summary', rows), 'Summary')
                    self.assertEqual(loaded[1:], data)

    def test_duplicate_metrics_and_missing_values_are_rejected(self):
        for case in SUMMARY_CASES:
            for rows in ([['Total_Courses', 22], [' total_courses ', 22]], [['Total_Courses']]):
                with self.subTest(case=case, rows=rows):
                    self.assertIsNone(grader(case).load_sheet_rows(book('Summary', rows), 'Summary'))

    def test_incorrect_numeric_value_is_preserved_for_existing_checks(self):
        for case in SUMMARY_CASES:
            with self.subTest(case=case):
                g = grader(case)
                rows = g.load_sheet_rows(book('Summary', [['Total_Courses', -500]]), 'Summary')
                self.assertFalse(g.num_close(rows[1][1], 22))


class SummaryCriticalGateTests(unittest.TestCase):
    def test_exam_invalid_summary_fails_critical_metric_gate(self):
        case = 'canvas-exam-prep-scheduler'
        for mutation in ('duplicate', 'missing-value', 'missing-sheet'):
            with self.subTest(mutation=mutation):
                g = grader(case)
                wb = book('Quiz Performance', [['Course', 'Quiz', 'Avg_Score', 'Below_Threshold'],
                                              ['Algebra', 'Quiz', 70, 'Yes']])
                ws = wb.create_sheet('Review Schedule')
                ws.append(['Course', 'Topic', 'Date', 'Time', 'Room'])
                ws.append(['Algebra', 'Quiz Review Algebra', '2026-03-16', '16:00', 'Room 101'])
                ws = wb.create_sheet('Summary')
                for row in [('Total_Quizzes_Analyzed', 1), ('Below_Threshold_Quizzes', 1),
                            ('Courses_Needing_Review', 1), ('Review_Sessions_Scheduled', 1)]:
                    ws.append(row)
                self.break_summary(wb, mutation)
                with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
                    wb.save(Path(tmp) / 'Exam_Prep.xlsx')
                    expected = [('Algebra', 'Quiz', 70)]
                    g.check_excel(tmp, expected, expected, {'Algebra'})
                critical = 'Summary: Total_Quizzes_Analyzed и Below_Threshold_Quizzes верны'
                self.assertIn(critical, g.FAILED_NAMES)
                self.assertIn(critical, g.CRITICAL_CHECKS)

    def test_ta_invalid_summary_fails_critical_metric_gate(self):
        case = 'canvas-ta-workload-excel-email'
        source = ROOT / 'tasks/finalpool' / case / 'groundtruth_workspace/TA_Workload_Report.xlsx'
        for mutation in ('duplicate', 'missing-value', 'missing-sheet'):
            with self.subTest(mutation=mutation):
                g = grader(case)
                wb = openpyxl.load_workbook(source)
                self.break_summary(wb, mutation)
                with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
                    wb.save(Path(tmp) / 'TA_Workload_Report.xlsx')
                    g.check_excel(tmp)
                self.assertIn('Total_Courses = 22', g.FAILED_NAMES)
                self.assertIn('Total_Courses = 22', g.CRITICAL_CHECKS)

    @staticmethod
    def break_summary(wb, mutation):
        ws = wb['Summary']
        row = 2 if ws.cell(1, 1).value == 'Metric' else 1
        if mutation == 'duplicate':
            ws.append([ws.cell(row, 1).value, 999])
        elif mutation == 'missing-value':
            ws.cell(row, 2).value = None
        else:
            del wb['Summary']


class SummaryVerdictTests(unittest.TestCase):
    CASE_FILES = {
        'canvas-grade-summary': 'Canvas_Grade_Summary.xlsx',
        'canvas-quiz-report': 'Canvas_Quiz_Report.xlsx',
        'sf-hr-manager-report': 'HR_Manager_Report.xlsx',
        'sf-hr-salary-growth': 'HR_Salary_Growth.xlsx',
    }

    def verdict(self, case, filename, mutate):
        g = grader(case)
        source = ROOT / 'tasks/finalpool' / case / 'groundtruth_workspace'
        wb = openpyxl.load_workbook(source / filename)
        ws = wb['Summary']
        ws.delete_rows(1)
        mutate(ws)
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            wb.save(Path(tmp) / filename)
            with patch.object(sys, 'argv', ['grader', '--agent_workspace', tmp, '--groundtruth_workspace', str(source)]):
                try:
                    g.main()
                except SystemExit as result:
                    return result.code
        return 0

    def test_full_offline_grader_accepts_headerless_summary(self):
        for case, filename in self.CASE_FILES.items():
            with self.subTest(case=case):
                self.assertEqual(self.verdict(case, filename, lambda ws: None), 0)

    def test_full_grader_rejects_wrong_missing_or_duplicate_metric(self):
        mutations = [lambda ws: setattr(ws.cell(1, 2), 'value', -999999),
                     lambda ws: ws.delete_rows(1),
                     lambda ws: ws.append([c.value for c in ws[1]])]
        for case, filename in self.CASE_FILES.items():
            for mutate in mutations:
                with self.subTest(case=case, mutation=mutate):
                    self.assertNotEqual(self.verdict(case, filename, mutate), 0)


class EquityTests(unittest.TestCase):
    CASE = 'sf-hr-compensation-equity-excel-word-forms'

    def run_excel(self, mutate):
        g = grader(self.CASE)
        source = ROOT / 'tasks/finalpool' / self.CASE / 'groundtruth_workspace'
        wb = openpyxl.load_workbook(source / 'Compensation_Equity.xlsx')
        mutate(wb['Equity Metrics'])
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            wb.save(Path(tmp) / 'Compensation_Equity.xlsx')
            g.check_excel(tmp, str(source))
        return g.CRITICAL_FAILS

    def test_extra_columns_and_reordered_required_columns(self):
        def mutate(ws):
            ws.cell(1, 2, 'Experience Band')
            ws.cell(1, 3, 'Highest Paid Education')
            ws.cell(1, 4, 'Lowest Paid Education')
            ws.cell(1, 5, 'Pay Gap %')
            ws.insert_cols(5, 2)
            ws.cell(1, 5, 'HighestAvgSalary')
            ws.cell(1, 6, 'LowestAvgSalary')
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 5, 61568.61)
                ws.cell(row, 6, 60047.22)
            for row in range(1, ws.max_row + 1):
                a, b = ws.cell(row, 7), ws.cell(row, 8)
                a.value, b.value = b.value, a.value
        self.assertEqual(self.run_excel(mutate), [])

    def test_wrong_values_missing_headers_duplicate_headers_and_rows_rejected(self):
        mutations = [
            lambda ws: setattr(ws.cell(2, 5), 'value', -999),
            lambda ws: setattr(ws.cell(1, 5), 'value', 'Missing'),
            lambda ws: setattr(ws.cell(1, 8), 'value', ws.cell(1, 5).value),
            lambda ws: ws.append([c.value for c in ws[2]]),
        ]
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                self.assertTrue(self.run_excel(mutate))


class ExamScheduleTests(unittest.TestCase):
    def check_schedule(self, date='2026-03-16', time='16:00-17:30'):
        g = grader('canvas-exam-prep-scheduler')
        wb = book('Review Schedule', [
            ['Course', 'Topic', 'Date', 'Time', 'Room'],
            ['Algebra', 'Quiz Review Algebra', date, time, 'Room 101'],
        ])
        seen = []
        g.check = lambda name, ok, detail='': seen.append((name, ok))
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            wb.save(Path(tmp) / 'Exam_Prep.xlsx')
            g.check_excel(tmp, [('Algebra', 'Quiz', 70)], [('Algebra', 'Quiz', 70)], {'Algebra'})
        return next(ok for name, ok in seen if name.startswith('Review Schedule:'))

    def test_task_column_order(self):
        self.assertTrue(self.check_schedule())

    def test_weekend_wrong_date_wrong_time_rejected(self):
        for date, time in [('2026-03-15', '16:00'), ('2026-03-14', '16:00'), ('2026-03-16', '09:00')]:
            self.assertFalse(self.check_schedule(date, time))


class CurriculumTests(unittest.TestCase):
    def check_c3(self, mutate):
        g = grader('terminal-canvas-sf-excel-ppt-gcal')
        wb = book('Curriculum_Coverage', [
            ['Curriculum coverage'],
            ['Course ID', 'Course Name', 'Enrolled Students', 'Assignments', 'Quizzes', 'Total_Assessments'],
            *[[i, f'Course {i}', 100, 10, 2, 12] for i in range(1, 5)],
        ])
        mutate(wb.active)
        seen = []
        g.critical = lambda name, ok, detail='': seen.append((name, ok, detail))
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            wb.save(Path(tmp) / 'Skills_Gap_Analysis.xlsx')
            g.critical_checks(tmp)
        return next(ok for name, ok, detail in seen if name.startswith('C3 '))

    def test_title_row_is_allowed(self):
        self.assertTrue(self.check_c3(lambda ws: None))

    def test_wrong_total_missing_column_and_duplicate_course_rejected(self):
        for mutate in [lambda ws: setattr(ws.cell(3, 6), 'value', 999),
                       lambda ws: setattr(ws.cell(2, 5), 'value', 'Missing'),
                       lambda ws: setattr(ws.cell(6, 1), 'value', 1)]:
            with self.subTest(mutate=mutate):
                self.assertFalse(self.check_c3(mutate))


if __name__ == '__main__':
    unittest.main()
