"""Native workbook-checker regressions; fixtures preserve archived output cell values."""
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(os.environ.get("GRADER_REPO_ROOT", Path(__file__).resolve().parents[2]))
CASES = {
    'arxiv': 'arxiv-research-tracker-teamly',
    'coupon': 'pw-insales-coupon-effectiveness-gsheet-email',
    'hr': 'sf-hr-attrition-forecast-excel-word-gcal',
    'canvas': 'terminal-canvas-gsheet-word-teamly-gcal',
    'youtube': 'yt-fireship-scholarly-excel-teamly',
}


def load(case):
    spec = importlib.util.spec_from_file_location('grader', ROOT / 'tasks/finalpool' / CASES[case] / 'evaluation/main.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class NativeCheckers(unittest.TestCase):
    def evaluate(self, case, mutate=lambda wb: None):
        module = load(case)
        fixture = json.loads((Path(__file__).parent / 'fixtures' / (CASES[case] + '.json')).read_text())
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in fixture['sheets'].items():
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append(row)
        mutate(wb)
        results = {}
        hook = 'check' if case in ('coupon', 'canvas') else 'record'
        original = getattr(module, hook)
        def record(name, condition, *args, **kwargs):
            results[name] = bool(condition)
            return original(name, condition, *args, **kwargs)
        with tempfile.TemporaryDirectory() as directory, contextlib.redirect_stdout(io.StringIO()), patch.object(module, hook, record):
            wb.save(Path(directory) / fixture['filename'])
            wb.close()
            if case == 'coupon':
                # External services are intentionally unavailable; only workbook checks are asserted.
                with patch.object(module, 'get_conn', side_effect=RuntimeError('offline workbook regression')):
                    try:
                        module.run_evaluation(directory, directory, None, None)
                    except SystemExit as exc:
                        self.assertEqual(exc.code, 1)
            elif case == 'hr':
                gt = module.get_groundtruth(str(ROOT / 'tasks/finalpool' / CASES[case] / 'groundtruth_workspace'))
                module.check_excel(directory, gt)
            elif case == 'canvas':
                # Fixed reference isolates header recognition from live database availability.
                reference = {20: {'avg': 80.26, 'passrate': 91, 'count': 100, 'failing': 9},
                             21: {'avg': 79.08, 'passrate': 90, 'count': 100, 'failing': 10},
                             'combined': {'count': 200, 'failing': 19}}
                with patch.object(module, 'compute_canvas_analytics', return_value=reference):
                    module.check_excel(directory)
            else:
                module.check_excel(directory)
        return results

    def relevant(self, case, results):
        if case == 'arxiv':
            return results['Paper Comparison contains exactly the 5 target papers (no noise)']
        if case == 'coupon':
            return all(results[n] for n in results if n.startswith(('Data_Analysis Category column carries', 'Data_Analysis normalized columns', 'Market_Avg_Price matches', 'Price_Gap_Pct =')))
        if case == 'hr':
            return all(results[n] for n in results if n.endswith(' Top_Factor'))
        if case == 'canvas':
            return results['Course_Summary averages match live Canvas (per course, tol 1.0)']
        return results['Video_Paper_Mapping has all 8 expected video titles in order']



    def test_arxiv_prefixed_id_still_validates_citations(self):
        def mutate(wb):
            ws = wb['Paper Comparison']
            column = [cell.value for cell in ws[1]].index('Citation_Count') + 1
            ws.cell(2, column).value = 999999
        results = self.evaluate('arxiv', mutate)
        self.assertTrue(self.relevant('arxiv', results))
        self.assertFalse(results['Citation counts exact for all 5 papers'])

    def test_coupon_alias_still_validates_market_value(self):
        def mutate(wb):
            ws = wb['Data_Analysis']
            column = [cell.value for cell in ws[1]].index('Benchmark_Market_Avg_Price') + 1
            ws.cell(2, column).value = 999999
        self.assertFalse(self.relevant('coupon', self.evaluate('coupon', mutate)))

    def coupon_duplicate_market(self, wrong_first=False, wrong_last=False, alias_first=False):
        def mutate(wb):
            ws = wb['Data_Analysis']
            column = [cell.value for cell in ws[1]].index('Benchmark_Market_Avg_Price') + 1
            values = [ws.cell(row, column).value for row in range(2, ws.max_row + 1)]
            ws.insert_cols(column)
            ws.cell(1, column).value = 'Benchmark_Market_Avg_Price' if alias_first else 'Market_Avg_Price'
            ws.cell(1, column + 1).value = 'Market_Avg_Price' if alias_first else 'Benchmark_Market_Avg_Price'
            for row, value in enumerate(values, 2):
                ws.cell(row, column).value = 999999 if wrong_first else value
                if wrong_last:
                    ws.cell(row, column + 1).value = 999999
        return self.evaluate('coupon', mutate)

    def test_coupon_conflicting_canonical_before_alias_rejected(self):
        self.assertFalse(self.relevant('coupon', self.coupon_duplicate_market(wrong_first=True)))

    def test_coupon_conflicting_alias_after_canonical_rejected(self):
        self.assertFalse(self.relevant('coupon', self.coupon_duplicate_market(wrong_last=True)))

    def test_coupon_conflicting_alias_before_canonical_rejected(self):
        self.assertFalse(self.relevant('coupon', self.coupon_duplicate_market(wrong_first=True, alias_first=True)))

    def test_coupon_conflicting_category_alias_rejected(self):
        def mutate(wb):
            ws = wb['Data_Analysis']
            ws.insert_cols(1)
            ws.cell(1, 1).value = 'Category'
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 1).value = 'Unrelated category'
        self.assertFalse(self.relevant('coupon', self.evaluate('coupon', mutate)))

    def test_coupon_equal_canonical_and_alias_accepted(self):
        self.assertTrue(self.relevant('coupon', self.coupon_duplicate_market()))

    def test_canvas_mean_alias_still_validates_numeric_value(self):
        def mutate(wb):
            wb['Course_Summary'].cell(2, 4).value = 0
        self.assertFalse(self.relevant('canvas', self.evaluate('canvas', mutate)))

    def test_hr_policy_field_alias(self):
        def mutate(wb):
            for row in wb['Top Risk Factors'].iter_rows(min_row=2):
                row[1].value = 'JOB_SATISFACTION'
        self.assertTrue(self.relevant('hr', self.evaluate('hr', mutate)))

    def test_youtube_curly_apostrophe_does_not_allow_wrong_order(self):
        def mutate(wb):
            ws = wb['Video_Paper_Mapping']
            ws.cell(2, 2).value, ws.cell(3, 2).value = ws.cell(3, 2).value, ws.cell(2, 2).value
        self.assertFalse(self.relevant('youtube', self.evaluate('youtube', mutate)))

    def test_hr_satisfaction_substring_is_not_alias(self):
        def mutate(wb):
            for row in wb['Top Risk Factors'].iter_rows(min_row=2):
                row[1].value = 'Not Job Satisfaction'
        self.assertFalse(self.relevant('hr', self.evaluate('hr', mutate)))


def change(case, wb, canonical=False):
    if case == 'arxiv':
        ws = wb['Paper Comparison']
        for row in ws.iter_rows(min_row=2):
            row[0].value = row[0].value.removeprefix('arXiv:')
        if not canonical:
            ws.cell(2, 1).value = 'arXiv:2402.99999'
    elif case == 'coupon':
        ws = wb['Data_Analysis']
        aliases = {'Product_Category':'Category', 'Internal_Product_Count':'Product_Count', 'Internal_Avg_Price':'Our_Avg_Price', 'Benchmark_Market_Avg_Price':'Market_Avg_Price'}
        for cell in ws[1]:
            cell.value = aliases.get(cell.value, cell.value)
        if not canonical:
            ws.cell(2, 1).value = 'Unrelated category'
    elif case == 'hr':
        ws = wb['Top Risk Factors']
        for row in ws.iter_rows(min_row=2):
            row[1].value = 'satisfaction' if canonical else 'performance'
    elif case == 'canvas':
        ws = wb['Course_Summary']
        for cell in ws[1]:
            if cell.value == 'mean_score':
                cell.value = 'average_score'
        if not canonical:
            ws.cell(2, 4).value = 0
    else:
        ws = wb['Video_Paper_Mapping']
        for row in ws.iter_rows(min_row=2):
            row[1].value = row[1].value.replace('’', "'")
        if not canonical:
            ws.cell(8, 2).value = "Microsoft's unrelated video"


def make_test(case, variant):
    def test(self):
        mutate = (lambda wb: None) if variant == 'archived_alias' else (lambda wb: change(case, wb, variant == 'canonical'))
        results = self.evaluate(case, mutate)
        self.assertEqual(self.relevant(case, results), variant != 'wrong_value', results)
    return test


for case in CASES:
    for variant in ('archived_alias', 'canonical', 'wrong_value'):
        setattr(NativeCheckers, f'test_{case}_{variant}', make_test(case, variant))

if __name__ == '__main__':
    unittest.main()
