"""Execute actual graders with frozen artifact fixtures and injected database rows."""

import ast
import contextlib
import copy
import io
import json
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest

from docx import Document
import openpyxl

ROOT = Path(os.environ.get("COWORK_GRADER_ROOT", Path(__file__).resolve().parents[2]))
FIXTURES = Path(__file__).with_name("fixtures") / "artifact_extraction"


class Database:
    def __init__(self, results):
        self.results = iter(results)
        self.rows = []
        self.queries = []

    def cursor(self):
        return self

    def execute(self, sql, args=None):
        self.queries.append((sql, args))
        if args is not None:
            # psycopg2 treats unescaped literal percentages as placeholders.
            sql % tuple("value" for _ in args)
        self.rows = next(self.results)
        if "HR_ANALYTICS__PUBLIC__EMPLOYEES" in sql and '"EMPLOYEE_ID"' not in sql:
            self.rows = [row[1:] if len(row) == 3 else row for row in self.rows]

    def fetchall(self):
        return self.rows

    def fetchone(self):
        return self.rows[0]

    def close(self):
        pass


def grader(case, results=()):
    path = ROOT / "tasks/finalpool" / case / "evaluation/main.py"
    tree = ast.parse(path.read_text())
    tree.body = [
        n
        for n in tree.body
        if not (
            isinstance(n, ast.Import) and any(a.name == "psycopg2" for a in n.names)
        )
    ]
    db = Database(results)
    env = {
        "__name__": "regression_grader",
        "__file__": str(path),
        "psycopg2": SimpleNamespace(connect=lambda **kwargs: db),
    }
    exec(compile(tree, str(path), "exec"), env)
    checks = {}

    def record(name, condition, detail="", **kwargs):
        checks[name] = (bool(condition), str(detail))

    env["check"] = env["record"] = record
    return env, checks, db


def invoke(env, name, *args):
    with contextlib.redirect_stdout(io.StringIO()):
        return env[name](*args)


def outcome(checks, fragment):
    selected = [v for k, v in checks.items() if fragment in k]
    if len(selected) != 1:
        raise AssertionError((fragment, checks))
    return selected[0][0]


class ArtifactExtraction(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.workspace = Path(self.tmp.name)

    def test_course_semesters_do_not_collide(self):
        expected = {
            1: {
                "name": "Chemistry (Fall 2014)",
                "avg_score": 60,
                "risk_level": "critical",
            },
            2: {
                "name": "Chemistry (Spring 2015)",
                "avg_score": 90,
                "risk_level": "on track",
            },
        }
        header = [
            "Course_Name",
            "Total_Enrollment",
            "Avg_Score",
            "Completion_Rate",
            "Risk_Level",
        ]
        for wrong in [False, True]:
            rows = [
                header,
                [expected[1]["name"], 10, 90 if wrong else 60, 80, "Critical"],
                [expected[2]["name"], 10, 30 if wrong else 90, 80, "On Track"],
            ]
            cells = [(i, j, v) for i, row in enumerate(rows) for j, v in enumerate(row)]
            env, checks, _ = grader(
                "course-enrollment-analytics-dashboard", [[("s", "Analytics")], cells]
            )
            invoke(env, "check_gsheet", expected)
            self.assertEqual(outcome(checks, "avg_score корректен"), not wrong, checks)
        for names, accepted in [
            ("Chemistry (Fall 2014)", True),
            ("Chemistry", False),
            ("Chemistry (Spring 2015)", False),
            ("Chemistry (Fall 2014) Chemistry (Spring 2015)", False),
        ]:
            env, checks, _ = grader(
                "course-enrollment-analytics-dashboard", [[(1, "Plan", names)]]
            )
            invoke(env, "check_teamly", expected, [expected[1]["name"]])
            self.assertEqual(outcome(checks, "ровно курсы"), accepted, checks)

    def test_archived_course_outputs_keep_semesters(self):
        data = json.loads((FIXTURES / "course_outputs.json").read_text())
        expected = {
            r["course_id"]: dict(
                r, name=r["course_name"], risk_level=r["risk_level"].lower()
            )
            for r in data["performance"]
            if r["avg_score"] is not None
        }
        cells = [
            (i, j, value)
            for i, row in enumerate(data["sheet"])
            for j, value in enumerate(row)
        ]
        env, checks, _ = grader(
            "course-enrollment-analytics-dashboard", [[("sheet", "Analytics")], cells]
        )
        invoke(env, "check_gsheet", expected)
        self.assertTrue(outcome(checks, "avg_score корректен"), checks)
        env, checks, _ = grader(
            "course-enrollment-analytics-dashboard", [[(1, "Plan", data["page"])]]
        )
        invoke(
            env,
            "check_teamly",
            expected,
            [r["name"] for r in expected.values() if r["risk_level"] != "on track"],
        )
        self.assertTrue(outcome(checks, "ровно курсы"), checks)

    def test_course_noise_query_keeps_literal_percentages_separate_from_parameters(
        self,
    ):
        env, checks, db = grader(
            "course-enrollment-analytics-dashboard", [[], [(0,)], [], [(1,)]]
        )
        invoke(env, "check_email_calendar", [])
        self.assertTrue(outcome(checks, "Шумовые письма"))
        self.assertEqual(len(db.queries), 4)

    def test_docx_table_products_and_nested_table_are_read(self):
        source = FIXTURES / "Root_Cause_Report.docx"
        doc = Document(source)
        products = [
            c.text
            for t in doc.tables
            for r in t.rows
            for c in r.cells
            if c.text.startswith(("MICROSMT", "AGARO"))
        ]
        self.assertEqual(len(products), 2)
        for mode in ["original", "nested", "missing"]:
            with self.subTest(mode=mode):
                doc = Document(source)
                if mode != "original":
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                if cell.text in products:
                                    text = cell.text
                                    cell.text = ""
                                    if mode == "nested":
                                        cell.add_table(rows=1, cols=1).cell(
                                            0, 0
                                        ).text = text
                doc.save(self.workspace / "Root_Cause_Report.docx")
                env, checks, _ = grader("insales-refund-root-cause-excel-word-email")
                invoke(env, "check_word", str(self.workspace), products)
                self.assertEqual(
                    outcome(checks, "mentions investigation products"),
                    mode != "missing",
                    checks,
                )

    def test_compliance_overall_status_comes_from_its_cell_not_summary_or_other_columns(
        self,
    ):
        page = json.loads((FIXTURES / "compliance_page.json").read_text())[0]
        lines = [line for line in page["body"].splitlines() if line.startswith("|")]
        records = [
            [s.strip() for s in line.strip("|").split("|")] for line in lines[2:]
        ]
        courses = [
            {
                "name": r[0],
                "short": r[0].split("(")[0].strip().lower(),
                "compliant": r[6] == "Compliant",
            }
            for r in records
        ]
        for mode in [
            "original",
            "html",
            "wrong_status",
            "wrong_noncompliant_status",
            "missing_course",
            "wrong_followup",
        ]:
            body = page["body"]
            if mode == "html":
                body = (
                    "<table>"
                    + "".join(
                        "<tr>" + "".join("<td>" + c + "</td>" for c in r) + "</tr>"
                        for r in [
                            [s.strip() for s in lines[0].strip("|").split("|")],
                            *records,
                        ]
                    )
                    + "</table>Non-Compliant summary"
                )
            if mode in ["wrong_status", "wrong_noncompliant_status"]:
                row = next(
                    r
                    for r in records
                    if r[6]
                    == ("Compliant" if mode == "wrong_status" else "Non-Compliant")
                )
                before = next(
                    line
                    for line in body.splitlines()
                    if line.startswith("|") and row[0] in line
                )
                changed = row.copy()
                changed[6] = "Non-Compliant" if mode == "wrong_status" else "Compliant"
                body = body.replace(before, "| " + " | ".join(changed) + " |")
            elif mode == "missing_course":
                body = "\n".join(
                    l for l in body.splitlines() if records[-1][0] not in l
                )
            elif mode == "wrong_followup":
                body = body.replace("2026-04-15", "2026-04-16")
            env, checks, _ = grader(
                "playwright-canvas-curriculum-word-teamly", [[(1, page["title"], body)]]
            )
            invoke(env, "check_teamly", courses)
            key = (
                "Follow_Up_Date present"
                if mode == "wrong_followup"
                else "Overall_Status correct"
            )
            self.assertEqual(outcome(checks, key), mode in ["original", "html"], checks)

    def test_exact_average_metric_does_not_collide_with_percentage(self):
        for mode in ["original", "reordered", "missing", "wrong"]:
            wb = openpyxl.load_workbook(FIXTURES / "Survey_Report.xlsx")
            ws = wb["Metrics"]
            row = next(r[0].row for r in ws.iter_rows() if r[0].value == "Avg_Gap")
            if mode == "missing":
                ws.delete_rows(row)
            elif mode == "wrong":
                ws.cell(row, 2, 1234)
            elif mode == "reordered":
                records = list(ws.values)
                ws.delete_rows(1, ws.max_row)
                for r in [records[0], *reversed(records[1:])]:
                    ws.append(r)
            wb.save(self.workspace / "Survey_Report.xlsx")
            wb.close()
            env, checks, db = grader("pw-teamly-forms-survey-excel-word", [[], []])
            env["get_conn"] = lambda: db
            invoke(env, "run_evaluation", str(self.workspace), "", "", None)
            self.assertEqual(
                outcome(checks, "Total_Items и Avg_Gap"),
                mode in ["original", "reordered"],
                checks,
            )

    def test_paper_identity_prefers_explicit_id_over_author_substrings(self):
        pages = json.loads((FIXTURES / "paper_pages.json").read_text())
        for mode in ["original", "missing_t5", "wrong_id"]:
            rows = []
            for p in pages:
                if "1910.10683" in p["body"]:
                    if mode == "missing_t5":
                        continue
                    if mode == "wrong_id":
                        p = dict(p, body=p["body"].replace("1910.10683", "9999.99999"))
                rows.append((p["title"], p["body"]))
            env, checks, _ = grader(
                "terminal-arxiv-scholarly-teamly-word-excel",
                [[(1, "RPT", "Research Paper Tracker")], rows],
            )
            invoke(env, "check_teamly")
            self.assertEqual(
                outcome(checks, "Teamly has >= 6"), mode == "original", checks
            )
            # Source overlap mismatch is independent: unchanged frozen output stays FAIL.
            self.assertFalse(outcome(checks, "correct Category/Source"), checks)

    def test_wrapped_menu_and_root_forms_preserve_actual_day_semantics(self):
        original = json.loads((FIXTURES / "evidence_based_menus.json").read_text())
        for mode in [
            "wrapped",
            "list",
            "dict",
            "duplicate_category",
            "missing_day",
            "missing_lunch",
        ]:
            data = copy.deepcopy(original)
            if mode == "duplicate_category":
                data["menus"][1]["lunch_category"] = data["menus"][0]["lunch_category"]
            if mode == "missing_day":
                data["menus"].pop()
            if mode == "missing_lunch":
                for r in data["menus"]:
                    r.pop("lunch", None)
                    r.pop("lunch_category", None)
            if mode == "list":
                data = data["menus"]
            if mode == "dict":
                data = {r["day"]: r for r in data["menus"]}
            (self.workspace / "evidence_based_menus.json").write_text(json.dumps(data))
            env, checks, _ = grader("terminal-kulinar-scholarly-excel-word-forms")
            invoke(env, "check_menu_rule", str(self.workspace))
            accepted = all(v[0] for v in checks.values())
            self.assertEqual(accepted, mode in ["wrapped", "list", "dict"], checks)

    def test_markdown_paper_id_remains_authoritative(self):
        pages = json.loads((FIXTURES / "paper_pages.json").read_text())
        for label in ["**Paper_ID:**", "**Paper_ID**:", "Paper_ID:", "`Paper_ID`:"]:
            for value in ["1910.10683", "9999.99999", "invalid", ""]:
                with self.subTest(label=label, value=value):
                    rows = []
                    for page in pages:
                        body = page["body"]
                        if "1910.10683" in body:
                            body = body.replace("**Paper_ID:** 1910.10683", label + " " + value)
                        rows.append((page["title"], body))
                    env, checks, _ = grader(
                        "terminal-arxiv-scholarly-teamly-word-excel",
                        [[(1, "RPT", "Research Paper Tracker")], rows],
                    )
                    invoke(env, "check_teamly")
                    self.assertEqual(outcome(checks, "Teamly has >= 6"), value == "1910.10683", checks)

    def test_menu_weekday_identity_and_chronological_adjacency(self):
        original = json.loads((FIXTURES / "evidence_based_menus.json").read_text())
        for shape in ["wrapped", "list", "map"]:
            for mode in ["valid", "shuffled", "duplicate_day", "missing_day_label", "missing_day", "adjacent_category"]:
                with self.subTest(shape=shape, mode=mode):
                    days = copy.deepcopy(original["menus"])
                    if mode == "duplicate_day":
                        for day in days:
                            day["day"] = "Monday"
                    if mode == "missing_day_label":
                        for day in days:
                            day.pop("day")
                    if mode == "missing_day":
                        days.pop()
                    if mode == "adjacent_category":
                        days[1]["lunch_category"] = days[0]["lunch_category"]
                    if mode in ["shuffled", "adjacent_category"]:
                        days = [days[i] for i in [0, 2, 4, 1, 3]]
                    data = {"menus": days} if shape == "wrapped" else days
                    if shape == "map":
                        data = {day.get("day", "unknown" + str(i)): day for i, day in enumerate(days)}
                    (self.workspace / "evidence_based_menus.json").write_text(json.dumps(data))
                    env, checks, _ = grader("terminal-kulinar-scholarly-excel-word-forms")
                    invoke(env, "check_menu_rule", str(self.workspace))
                    self.assertEqual(all(v[0] for v in checks.values()), mode in ["valid", "shuffled"], checks)

    def test_bonus_rows_keys_and_duplicate_names_keep_identity(self):
        env, checks, _ = grader("terminal-moex-sf-gsheet-word-gcal")
        employees = [
            {
                "employee_id": str(i),
                "name": "Same Name",
                "salary": 1000 * (i + 1),
                "region": "Region " + str(i),
                "bonus_percentage": 10,
                "bonus_amount": 100 * (i + 1),
                "adjusted_bonus": 100 * (i + 1),
                "market_factor": 1,
            }
            for i in range(3)
        ]
        self.assertEqual(
            env["_entries"](
                {"market_rows": [{"ticker": "ABC"}], "employees": employees}
            ),
            employees,
        )
        self.assertEqual(env["_get"](employees[0], "bonus amount", "bonus"), 100)
        self.assertIsNone(
            env["_get"]({"bonus_percentage": 10}, "bonus amount", "bonus")
        )

    def test_bonus_expected_source_does_not_overwrite_equal_names(self):
        env, _, _ = grader(
            "terminal-moex-sf-gsheet-word-gcal",
            [
                [("R" + str(i), 100) for i in range(5)],
                [(str(i), "Same Name", 1000 * (i + 1)) for i in range(3)],
            ],
        )
        env["load_tier_table"] = lambda: [(0, 1000, 10)]
        result = invoke(env, "compute_expected_bonuses")
        self.assertEqual(len(result), 3, result)

    def test_bonus_actual_critical_checks_reject_wrong_missing_and_duplicate_employees(
        self,
    ):
        source = [(str(i), "Same Name", 1000 * (i + 1)) for i in range(3)]
        expected_env, _, _ = grader(
            "terminal-moex-sf-gsheet-word-gcal",
            [[("R" + str(i), 100 + i * 100) for i in range(5)], source],
        )
        expected_env["load_tier_table"] = lambda: [(0, 200, 5), (200, 1000, 10)]
        expected = invoke(expected_env, "compute_expected_bonuses")
        employees = [
            {
                "employee_id": str(i) + ".0",
                "name": "Same Name",
                "salary": 1000 * (i + 1),
                "region": "R" + str(i),
                "bonus_percentage": 5 if i == 0 else 10,
                "bonus_amount": 50 if i == 0 else 100 * (i + 1),
                "adjusted_bonus": 50 if i == 0 else 100 * (i + 1),
                "market_factor": 1,
            }
            for i in range(3)
        ]
        for mode in [
            "valid",
            "tied_permutation",
            "unique_salary_fallback",
            "wrong_bonus",
            "missing_bonus",
            "missing_employee",
            "duplicate_id",
            "wrong_salary",
            "wrong_region",
            "unknown_id",
            "over_person_cap",
            "over_total_cap",
        ]:
            with self.subTest(mode=mode):
                current = copy.deepcopy(employees)
                adjusted = copy.deepcopy(employees)
                if mode == "tied_permutation":
                    current[0]["region"], current[1]["region"] = (
                        current[1]["region"],
                        current[0]["region"],
                    )
                    current[0]["bonus_amount"], current[1]["bonus_amount"] = 100, 100
                if mode == "unique_salary_fallback":
                    for row in current:
                        row.pop("employee_id")
                if mode == "wrong_bonus":
                    current[0]["bonus_amount"] = 900
                if mode == "missing_bonus":
                    current[0].pop("bonus_amount")
                if mode == "missing_employee":
                    current.pop()
                if mode == "duplicate_id":
                    current[1]["employee_id"] = current[0]["employee_id"]
                if mode == "wrong_salary":
                    current[0]["salary"] = 9999
                if mode == "wrong_region":
                    current[0]["region"] = "R4"
                if mode == "unknown_id":
                    current[0]["employee_id"] = "unknown"
                if mode == "over_person_cap":
                    adjusted[0]["adjusted_bonus"] = 300
                if mode == "over_total_cap":
                    for row in adjusted:
                        row["adjusted_bonus"] = 0.2 * row["salary"]
                (self.workspace / "current_bonuses.json").write_text(
                    json.dumps(current)
                )
                (self.workspace / "market_adjusted_bonuses.json").write_text(
                    json.dumps(
                        {"market_rows": [{"ticker": "ABC"}], "employees": adjusted}
                    )
                )
                env, checks, _ = grader("terminal-moex-sf-gsheet-word-gcal")
                env["BUDGET_CAP"] = 1000
                invoke(
                    env,
                    "critical_checks",
                    str(self.workspace),
                    (1, 0, 100, 100),
                    expected,
                )
                self.assertEqual(
                    all(ok for ok, _ in checks.values()),
                    mode in ["valid", "tied_permutation", "unique_salary_fallback"],
                    checks,
                )


if __name__ == "__main__":
    unittest.main()
