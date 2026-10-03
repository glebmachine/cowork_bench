"""The actual DOCX grader must bind the count to meeting, not missing, target."""

import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

from docx import Document

ROOT = Path(os.environ.get("COWORK_GRADER_ROOT", Path(__file__).resolve().parents[2]))
FIXTURES = Path(__file__).parent / "fixtures/segment_summary"


def verdict(text=None):
    spec = importlib.util.spec_from_file_location("segment_grader", ROOT / "tasks/finalpool/sf-sales-segment-word/evaluation/main.py")
    module = importlib.util.module_from_spec(spec)
    with patch.dict(sys.modules, {"psycopg2": types.ModuleType("psycopg2")}):
        spec.loader.exec_module(module)
    seen = {}
    module.check = lambda name, ok, detail="", **kwargs: seen.__setitem__(name, bool(ok))
    document = Document(FIXTURES / "Q4_Segment_Report.docx")
    db = {row.cells[0].text: {"actual": float(row.cells[2].text), "count": int(row.cells[3].text)} for row in document.tables[0].rows[1:]}
    if text is not None:
        document.paragraphs[1].text = "Total target 760000, actual 393653.23. " + text
    with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
        document.save(Path(tmp) / "Q4_Segment_Report.docx")
        module.check_word_doc(tmp, db)
    return seen


class SegmentSummaryTests(unittest.TestCase):
    def test_archived_document_passes_the_complete_word_check(self):
        self.assertTrue(all(verdict().values()), verdict())

    def test_equivalent_count_before_and_after_target_statement(self):
        for text in [
            "Количество сегментов, выполнивших или перевыполнивших свои планы: 0.",
            "Число сегментов, выполнивших план: 0.",
            "0 сегментов выполнили план.",
            "Ноль сегментов выполнили свои планы.",
            "Ни один из 4 сегментов не выполнил план.",
            "0 segments met their targets.",
            "Segments meeting or exceeding target: 0.",
        ]:
            with self.subTest(text=text):
                checks = verdict(text)
                self.assertTrue(checks["Summary states 0 segments meeting target"], checks)

    def test_wrong_opposite_unrelated_and_conflicting_counts_fail(self):
        for text in [
            "Количество сегментов, выполнивших свои планы: 2.",
            "Количество сегментов, выполнивших свои планы: 10.",
            "Количество сегментов, не выполнивших свои планы: 0.",
            "10 segments met their targets.",
            "0 segments missed their targets.",
            "0 сегментов не выполнили план.",
            "0 orders were cancelled. All segments met their targets.",
            "Количество сегментов, выполнивших план: 0. 2 segments met their targets.",
        ]:
            with self.subTest(text=text):
                self.assertFalse(verdict(text)["Summary states 0 segments meeting target"])

    def test_archived_fixture_hash(self):
        expected = json.loads((FIXTURES / "provenance.json").read_text())
        self.assertEqual(hashlib.sha256((FIXTURES / "Q4_Segment_Report.docx").read_bytes()).hexdigest(), expected["sha256"])


if __name__ == "__main__":
    unittest.main()
