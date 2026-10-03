"""Actual Canvas grader functions, archived outputs and paired negative controls."""

import contextlib
import hashlib
from datetime import datetime, timedelta
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch
from zoneinfo import ZoneInfo

import openpyxl

ROOT = Path(os.environ.get("COWORK_GRADER_ROOT", Path(__file__).resolve().parents[2]))
FIXTURES = Path(__file__).parent / "fixtures/canvas_language"


class Database:
    def __init__(self, rows):
        self.rows = rows

    def cursor(self):
        return self

    def execute(self, *args):
        pass

    def fetchall(self):
        return self.rows

    def close(self):
        pass


def load(case, rows=()):
    spec = importlib.util.spec_from_file_location(case, ROOT / "tasks/finalpool" / case / "evaluation/main.py")
    module = importlib.util.module_from_spec(spec)
    stub = types.ModuleType("psycopg2")
    stub.connect = lambda **kwargs: Database(rows)
    with patch.dict(sys.modules, {"psycopg2": stub}):
        spec.loader.exec_module(module)
    checks = {}
    module.check = module.record = lambda name, ok, detail="", **kwargs: checks.__setitem__(name, bool(ok))
    return module, checks


def fixture(case):
    return json.loads((FIXTURES / (case + ".json")).read_text())


def invoke(module, method, *args):
    with contextlib.redirect_stdout(io.StringIO()):
        getattr(module, method)(*args)


class CanvasLanguageTests(unittest.TestCase):
    def test_frozen_fixture_hashes_match_provenance(self):
        provenance = json.loads((FIXTURES / "provenance.json").read_text())
        for source in provenance["sources"]:
            self.assertEqual(hashlib.sha256((FIXTURES / source["fixture"]).read_bytes()).hexdigest(), source["sha256"])

    def calendar_verdict(self, summary=None, delta_days=0, duration=45, description=""):
        case = "canvas-grade-equity-excel-word-gcal"
        event = fixture(case)["arguments"]
        start = datetime.fromisoformat(event["start"]["dateTime"]).replace(
            tzinfo=ZoneInfo(event["start"]["timeZone"])
        ).astimezone(ZoneInfo("UTC")) + timedelta(days=delta_days)
        module, checks = load(case, [(summary or event["summary"], description, start, start + timedelta(minutes=duration))])
        invoke(module, "check_calendar")
        return next(ok for name, ok in checks.items() if name.startswith("CRITICAL: GCal"))

    def test_archived_russian_course_and_english_alias_are_both_valid(self):
        for title in [None, "Biochemistry and Bioinformatics Grade Equity Review"]:
            with self.subTest(title=title):
                self.assertTrue(self.calendar_verdict(title))

    def test_wrong_course_date_duration_or_review_is_not_rescued_by_aliases(self):
        for options in [
            {"summary": "Алгебра Grade Equity Review"},
            {"summary": "Алгебра Grade Equity Review", "description": "Биохимия и биоинформатика"},
            {"summary": "Биохимия и биоинформатика Lecture"},
            {"summary": "BiochemistryX Grade Equity Review"},
            {"delta_days": 10},
            {"duration": 20},
        ]:
            with self.subTest(options=options):
                self.assertFalse(self.calendar_verdict(**options))

    def module_verdict(self, body, expected=None):
        case = "canvas-module-completion-teamly-gcal"
        data = fixture(case)
        module, checks = load(case, [(1, data["arguments"]["title"], body)])
        invoke(module, "check_teamly", expected or data["source_module_counts"])
        return checks["Item_Count values match live Canvas module item counts"]

    def test_archived_english_module_rows_match_live_russian_names(self):
        data = fixture("canvas-module-completion-teamly-gcal")
        self.assertTrue(self.module_verdict(data["arguments"]["body"]))

    def test_module_names_match_in_both_directions_and_same_language(self):
        data = fixture("canvas-module-completion-teamly-gcal")
        english = {"Introduction": 15, "Week 1": 2, "Week 3": 1, "Week 5": 1, "Week 8": 1}
        russian = data["source_module_counts"]
        for language, expected in [("en", english), ("ru", russian)]:
            for body_language in ["en", "ru"]:
                body = data["arguments"]["body"]
                if body_language == "ru":
                    for en, ru in zip(english, russian):
                        body = body.replace(en, ru)
                with self.subTest(expected=language, body=body_language):
                    self.assertTrue(self.module_verdict(body, expected))

    def test_wrong_missing_duplicate_module_count_and_other_column_numbers_fail(self):
        data = fixture("canvas-module-completion-teamly-gcal")
        body = data["arguments"]["body"]
        expected = {"Introduction": 15, "Week 1": 2, "Week 3": 1, "Week 5": 1, "Week 8": 1}
        line = "| Week 1 | 2 | 2 Files | Not Started |"
        controls = [
            body.replace(line, "| Week 1 | 999 | 2 Files | Not Started |"),
            body.replace(line, "| Week 1 | 2 items? | 2 Files | Not Started |"),
            body.replace(line, "| Week 10 | 2 | 2 Files | Not Started |"),
            body.replace(line, ""),
            body + "\n" + line,
            body + "\n| Неделя 1 | 999 | 2 Files | Not Started |",
        ]
        for index, text in enumerate(controls):
            with self.subTest(index=index):
                self.assertFalse(self.module_verdict(text, expected))

    def test_module_counts_follow_named_columns_in_markdown_and_html(self):
        body = fixture("canvas-module-completion-teamly-gcal")["arguments"]["body"]
        rows = [[v.strip() for v in line.strip("|").split("|")] for line in body.splitlines() if line.startswith("|")]
        rows = [rows[0], *rows[2:]]
        rows = [[row[i] for i in [2, 0, 3, 1]] for row in rows]
        markdown = "\n".join("| " + " | ".join(row) + " |" for row in rows)
        html = "<table>" + "".join("<tr>" + "".join("<td>" + cell + "</td>" for cell in row) + "</tr>" for row in rows) + "</table>"
        for body in [markdown, html]:
            self.assertTrue(self.module_verdict(body))

    def test_unrelated_html_table_does_not_hide_markdown_tracker(self):
        body = fixture("canvas-module-completion-teamly-gcal")["arguments"]["body"]
        notes = "<table><tr><td>Notes</td><td>No blockers</td></tr></table>"
        for mixed in [notes + "\n" + body, body + "\n" + notes]:
            with self.subTest(mixed=mixed):
                self.assertTrue(self.module_verdict(mixed))

    def test_contradictory_html_module_count_does_not_hide_behind_markdown(self):
        body = fixture("canvas-module-completion-teamly-gcal")["arguments"]["body"]
        conflicting = "<table><tr><th>Module_Name</th><th>Item_Count</th></tr><tr><td>Week 1</td><td>999</td></tr></table>"
        self.assertFalse(self.module_verdict(body + "\n" + conflicting))

    def summary_verdict(self, change=None):
        workbook = openpyxl.load_workbook(FIXTURES / "Curriculum_Review.xlsx")
        if change:
            change(workbook["Summary"])
        with tempfile.TemporaryDirectory() as tmp:
            workbook.save(Path(tmp) / "Curriculum_Review.xlsx")
            workbook.close()
            module, checks = load("canvas-scholarly-curriculum-review")
            invoke(module, "check_excel", tmp)
        return checks["Summary math consistent (total ~22, compliant+non-compliant==total, rate==round)"]

    def test_archived_russian_summary_and_english_alias(self):
        self.assertTrue(self.summary_verdict())
        for label in ["Total Courses", "Total number of courses", "Всего курсов"]:
            self.assertTrue(self.summary_verdict(lambda ws: setattr(ws["A2"], "value", label)))

    def test_summary_wrong_totals_rates_and_faculty_decoys_fail(self):
        for cell, value in [("B2", 90), ("B3", 17), ("B5", "12%"), ("A2", "Общее количество преподавателей по всем курсам")]:
            with self.subTest(cell=cell, value=value):
                self.assertFalse(self.summary_verdict(lambda ws: setattr(ws[cell], "value", value)))
        self.assertFalse(self.summary_verdict(lambda ws: ws.append(["Total courses", 24])))


if __name__ == "__main__":
    unittest.main()
